"""Report export module: support PDF and Excel format export."""

import logging
import re
from pathlib import Path
from typing import Optional, Dict, Any
import pandas as pd
import numpy as np

logger = logging.getLogger(__name__)


def export_to_excel(
    output_path: Path,
    coaching_report_path: Optional[Path] = None,
    top_issues_path: Optional[Path] = None,
    events_path: Optional[Path] = None,
    corner_cards_path: Optional[Path] = None,
    stability_report_path: Optional[Path] = None,
    events_summary_path: Optional[Path] = None
) -> Path:
    """
    Export report to Excel file (multi-sheet format).
    
    Args:
        output_path: Output Excel file path
        coaching_report_path: Coaching advice report Markdown path
        top_issues_path: Top Issues CSV path
        events_path: Events CSV path
        corner_cards_path: Corner cards CSV path
        stability_report_path: Stability report CSV path
        events_summary_path: Events summary Markdown path
    
    Returns:
        Output file path
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("Exporting Excel requires installation of openpyxl: pip install openpyxl")
        raise ImportError("Please install openpyxl: pip install openpyxl")
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        # Sheet 1: Top Issues
        if top_issues_path and top_issues_path.exists():
            top_issues_df = pd.read_csv(top_issues_path)
            top_issues_df.to_excel(writer, sheet_name='Top Issues', index=False)
            logger.info(f"Added Sheet: Top Issues ({len(top_issues_df)} rows)")
        
        # Sheet 2: Events
        if events_path and events_path.exists():
            events_df = pd.read_csv(events_path)
            events_df.to_excel(writer, sheet_name='Events', index=False)
            logger.info(f"Added Sheet: Events ({len(events_df)} rows)")
        
        # Sheet 3: Corner Cards
        if corner_cards_path and corner_cards_path.exists():
            corner_cards_df = pd.read_csv(corner_cards_path)
            corner_cards_df.to_excel(writer, sheet_name='Corner Cards', index=False)
            logger.info(f"Added Sheet: Corner Cards ({len(corner_cards_df)} rows)")
        
        # Sheet 4: Stability Report
        if stability_report_path and stability_report_path.exists():
            stability_df = pd.read_csv(stability_report_path)
            stability_df.to_excel(writer, sheet_name='Stability Report', index=False)
            logger.info(f"Added Sheet: Stability Report ({len(stability_df)} rows)")
        
        # Sheet 5: Summary (text content)
        if coaching_report_path and coaching_report_path.exists():
            summary_text = coaching_report_path.read_text(encoding='utf-8')
            # Convert Markdown text to DataFrame (one cell per line)
            summary_lines = summary_text.split('\n')
            summary_df = pd.DataFrame({'Content': summary_lines})
            summary_df.to_excel(writer, sheet_name='Coaching Report', index=False)
            logger.info(f"Added Sheet: Coaching Report ({len(summary_lines)} rows)")
        
        # Format Excel
        workbook = writer.book
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            
            # Setup header row style
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Auto-adjust column width
            for column in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
    
    logger.info(f"Excel report exported: {output_path}")
    return output_path


def export_to_pdf(
    output_path: Path,
    coaching_report_path: Optional[Path] = None,
    top_issues_path: Optional[Path] = None,
    events_summary_path: Optional[Path] = None
) -> Path:
    """
    Export report to PDF file.
    
    Args:
        output_path: Output PDF file path
        coaching_report_path: Coaching advice report Markdown path
        top_issues_path: Top Issues CSV path
        events_summary_path: Events summary Markdown path
    
    Returns:
        Output file path
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        logger.error("Exporting PDF requires installation of reportlab: pip install reportlab")
        raise ImportError("Please install reportlab: pip install reportlab")
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Create PDF document
    doc = SimpleDocTemplate(str(output_path), pagesize=A4)
    story = []
    
    # Setup styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=30,
        alignment=1  # Center alignment
    )
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=12,
        spaceBefore=12
    )
    
    # Title
    story.append(Paragraph("Intelligent Driving Analysis Report", title_style))
    story.append(Spacer(1, 0.3*inch))
    
    # 1. Top Issues table
    if top_issues_path and top_issues_path.exists():
        story.append(Paragraph("Main Issue Analysis", heading_style))
        
        top_issues_df = pd.read_csv(top_issues_path)
        # Select key columns
        display_cols = ['turn', 'composite_score', 'score', 'stability_score', 'events_count', 'top_types']
        available_cols = [col for col in display_cols if col in top_issues_df.columns]
        
        if available_cols:
            # Prepare table data
            table_data = [['Turn', 'Composite Score', 'Issue Score', 'Stability Score', 'Event Count', 'Main Issues']]
            
            for _, row in top_issues_df.head(10).iterrows():
                row_data = []
                for col in available_cols:
                    val = row[col]
                    if pd.isna(val):
                        row_data.append('-')
                    elif isinstance(val, (int, np.integer)):
                        row_data.append(str(int(val)))
                    elif isinstance(val, float):
                        row_data.append(f"{val:.2f}")
                    else:
                        row_data.append(str(val))
                table_data.append(row_data)
            
            # Createtable
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
            ]))
            
            story.append(table)
            story.append(Spacer(1, 0.2*inch))
    
    # 2. Coaching advice report
    if coaching_report_path and coaching_report_path.exists():
        story.append(PageBreak())
        story.append(Paragraph("Coaching Advice Report", heading_style))
        
        # Read Markdown content and convert to PDF format
        markdown_text = coaching_report_path.read_text(encoding='utf-8')
        
        # Improved Markdown to PDF conversion
        lines = markdown_text.split('\n')
        table_data = []
        in_table = False
        table_headers = []
        
        for i, line in enumerate(lines):
            original_line = line
            line = line.strip()
            
            # Skip horizontal rules
            if line == '---' or line == '***':
                story.append(Spacer(1, 0.15*inch))
                continue
            
            if not line:
                if not in_table:
                    story.append(Spacer(1, 0.08*inch))
                continue
            
            # Process markdown tables
            if line.startswith('|') and '|' in line[1:]:
                if not in_table:
                    in_table = True
                    table_data = []
                
                # Parse table row
                cells = [cell.strip() for cell in line.split('|')[1:-1]]
                
                # Check if it's a separator row (contains ---)
                if any('---' in cell or ':' in cell for cell in cells):
                    continue
                
                table_data.append(cells)
                
                # If this is the first data row, it might be headers
                if len(table_data) == 1:
                    table_headers = cells
                continue
            else:
                # If we were in a table, render it now
                if in_table and len(table_data) > 0:
                    # Create table
                    table = Table(table_data)
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 11),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                        ('TOPPADDING', (0, 0), (-1, 0), 10),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('FONTSIZE', (0, 1), (-1, -1), 10),
                        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
                    ]))
                    story.append(table)
                    story.append(Spacer(1, 0.15*inch))
                    table_data = []
                    in_table = False
            
            # Process headings
            if line.startswith('# '):
                story.append(Spacer(1, 0.2*inch))
                text = line[2:].strip()
                # Remove emojis for better PDF compatibility
                text = text.replace('📊', '').replace('🎯', '').replace('📈', '').replace('🟢', '').replace('🟡', '').replace('🔴', '').strip()
                story.append(Paragraph(text, title_style))
            elif line.startswith('## '):
                story.append(Spacer(1, 0.15*inch))
                text = line[3:].strip()
                text = text.replace('📊', '').replace('🎯', '').replace('📈', '').replace('🟢', '').replace('🟡', '').replace('🔴', '').strip()
                story.append(Paragraph(text, heading_style))
            elif line.startswith('### '):
                story.append(Spacer(1, 0.12*inch))
                text = line[4:].strip()
                text = text.replace('📊', '').replace('🎯', '').replace('📈', '').replace('🟢', '').replace('🟡', '').replace('🔴', '').strip()
                story.append(Paragraph(text, styles['Heading3']))
            elif line.startswith('#### '):
                story.append(Spacer(1, 0.1*inch))
                text = line[5:].strip()
                text = text.replace('📊', '').replace('🎯', '').replace('📈', '').replace('🟢', '').replace('🟡', '').replace('🔴', '').strip()
                story.append(Paragraph(text, styles['Heading4']))
            elif line.startswith('- ') or line.startswith('* '):
                # List item - handle bold text
                text = line[2:].strip()
                # Convert **bold** to <b>bold</b>
                text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
                story.append(Paragraph(f"• {text}", styles['Normal']))
            else:
                # Regular paragraph - handle bold text
                text = line
                # Convert **bold** to <b>bold</b>
                text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
                # Remove emojis
                text = text.replace('📊', '').replace('🎯', '').replace('📈', '').replace('🟢', '').replace('🟡', '').replace('🔴', '').strip()
                if text:
                    story.append(Paragraph(text, styles['Normal']))
            
            story.append(Spacer(1, 0.05*inch))
        
        # Handle any remaining table
        if in_table and len(table_data) > 0:
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 11),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                ('TOPPADDING', (0, 0), (-1, 0), 10),
                ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('FONTSIZE', (0, 1), (-1, -1), 10),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.15*inch))
    
    # Generate PDF
    doc.build(story)
    logger.info(f"PDF report exported: {output_path}")
    return output_path


def markdown_to_pdf(markdown_path: Path, pdf_path: Path) -> Path:
    """
    Convert a Markdown file to a well-formatted PDF.
    
    Args:
        markdown_path: Input Markdown file path
        pdf_path: Output PDF file path
    
    Returns:
        Output PDF file path
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        logger.error("Exporting PDF requires installation of reportlab: pip install reportlab")
        raise ImportError("Please install reportlab: pip install reportlab")
    
    if not markdown_path.exists():
        raise FileNotFoundError(f"Markdown file not found: {markdown_path}")
    
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Create PDF document
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4, 
                           rightMargin=0.75*inch, leftMargin=0.75*inch,
                           topMargin=0.75*inch, bottomMargin=0.75*inch)
    story = []
    
    # Setup styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=20,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=20,
        spaceBefore=10,
        alignment=0  # Left alignment
    )
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=12,
        spaceBefore=16
    )
    subheading_style = ParagraphStyle(
        'CustomSubHeading',
        parent=styles['Heading3'],
        fontSize=14,
        textColor=colors.HexColor('#34495e'),
        spaceAfter=10,
        spaceBefore=12
    )
    
    # Read Markdown content
    markdown_text = markdown_path.read_text(encoding='utf-8')
    lines = markdown_text.split('\n')
    
    table_data = []
    in_table = False
    
    for i, line in enumerate(lines):
        original_line = line
        line = line.strip()
        
        # Skip horizontal rules
        if line == '---' or line == '***':
            story.append(Spacer(1, 0.15*inch))
            continue
        
        if not line:
            if not in_table:
                story.append(Spacer(1, 0.08*inch))
            continue
        
        # Process markdown tables
        if line.startswith('|') and '|' in line[1:]:
            if not in_table:
                in_table = True
                table_data = []
            
            # Parse table row
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            
            # Check if it's a separator row (contains --- or :)
            if any('---' in cell or (':' in cell and len(cell) < 10) for cell in cells):
                continue
            
            table_data.append(cells)
            continue
        else:
            # If we were in a table, render it now
            if in_table and len(table_data) > 0:
                # Create table
                table = Table(table_data)
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 11),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                    ('TOPPADDING', (0, 0), (-1, 0), 10),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                    ('FONTSIZE', (0, 1), (-1, -1), 10),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
                ]))
                story.append(table)
                story.append(Spacer(1, 0.15*inch))
                table_data = []
                in_table = False
        
        # Process headings
        if line.startswith('# '):
            story.append(Spacer(1, 0.2*inch))
            text = line[2:].strip()
            # Remove emojis for better PDF compatibility
            text = re.sub(r'[📊🎯📈🟢🟡🔴]', '', text).strip()
            story.append(Paragraph(text, title_style))
        elif line.startswith('## '):
            story.append(Spacer(1, 0.15*inch))
            text = line[3:].strip()
            text = re.sub(r'[📊🎯📈🟢🟡🔴]', '', text).strip()
            story.append(Paragraph(text, heading_style))
        elif line.startswith('### '):
            story.append(Spacer(1, 0.12*inch))
            text = line[4:].strip()
            text = re.sub(r'[📊🎯📈🟢🟡🔴]', '', text).strip()
            story.append(Paragraph(text, subheading_style))
        elif line.startswith('#### '):
            story.append(Spacer(1, 0.1*inch))
            text = line[5:].strip()
            text = re.sub(r'[📊🎯📈🟢🟡🔴]', '', text).strip()
            story.append(Paragraph(text, styles['Heading4']))
        elif line.startswith('- ') or line.startswith('* '):
            # List item - handle bold text
            text = line[2:].strip()
            # Convert **bold** to <b>bold</b>
            text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
            story.append(Paragraph(f"• {text}", styles['Normal']))
        else:
            # Regular paragraph - handle bold text
            text = line
            # Convert **bold** to <b>bold</b>
            text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
            # Remove emojis
            text = re.sub(r'[📊🎯📈🟢🟡🔴]', '', text).strip()
            if text:
                story.append(Paragraph(text, styles['Normal']))
        
        story.append(Spacer(1, 0.05*inch))
    
    # Handle any remaining table
    if in_table and len(table_data) > 0:
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('TOPPADDING', (0, 0), (-1, 0), 10),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.15*inch))
    
    # Generate PDF
    doc.build(story)
    logger.info(f"Markdown to PDF conversion completed: {pdf_path}")
    return pdf_path

